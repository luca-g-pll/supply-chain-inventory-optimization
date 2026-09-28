"""Input adapter: turn any daily sales file into the canonical long format.

Pipeline: ``read_any`` -> ``guess_mapping`` / ``ColumnMapping`` -> ``standardise``
-> ``validate``. The Kaggle demo goes through the same path (its ``sales``
column is mapped to ``quantity``), so the adapter is exercised on every run.

Contract (docs/07_data_dictionary.md §0): ``date``, ``item`` and ``quantity`` are
required; ``store``, ``unit_price``, ``unit_cost``, ``channel`` and
``lead_time_days`` are optional and fall back to documented defaults that are
flagged as *assumed*.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import BinaryIO

import numpy as np
import pandas as pd

from src.config import InputConfig, InventoryParams, PricesConfig
from src.series import DemandMatrix, natural_sorted

REQUIRED = ("date", "item", "quantity")
OPTIONAL = ("store", "unit_price", "unit_cost", "channel", "lead_time_days")
FIELDS = REQUIRED + OPTIONAL
SINGLE_STORE = "ALL"

# Column-name synonyms (EN / IT and a few common ERP names), already normalised.
SYNONYMS: dict[str, list[str]] = {
    "date": ["date", "data", "day", "giorno", "datetime", "timestamp", "order_date", "sale_date",
             "sales_date", "data_vendita", "data_ordine", "data_documento", "dt", "fecha", "datum"],
    "item": ["item", "sku", "product", "prodotto", "articolo", "codice", "code", "product_id", "item_id",
             "sku_id", "codice_articolo", "article", "material", "materiale", "item_code",
             "product_code", "art", "cod_articolo"],
    "quantity": ["quantity", "qty", "qta", "quantita", "units", "unit_sold", "units_sold", "sales",
                 "vendite", "venduto", "pezzi", "pz", "sold", "demand", "volume", "qty_sold",
                 "quantita_venduta", "sales_qty", "q"],
    "store": ["store", "negozio", "shop", "location", "punto_vendita", "pv", "filiale", "store_id",
              "branch", "warehouse", "magazzino", "sede", "site", "outlet", "store_code"],
    "unit_price": ["unit_price", "price", "prezzo", "prezzo_unitario", "selling_price",
                   "prezzo_vendita", "list_price", "retail_price", "sales_price"],
    "unit_cost": ["unit_cost", "cost", "costo", "costo_unitario", "purchase_price", "prezzo_acquisto",
                  "cogs", "buy_price", "costo_acquisto"],
    "channel": ["channel", "canale", "sales_channel", "canale_vendita"],
    "lead_time_days": ["lead_time_days", "lead_time", "leadtime", "lt", "tempo_consegna",
                       "giorni_consegna", "lead_days", "lead_time_giorni"],
}
_ONLINE = {"online", "web", "ecommerce", "e_commerce", "internet", "digital", "eshop", "e_shop",
           "webshop", "sito", "marketplace", "on_line"}
_PHYSICAL = {"physical", "store", "negozio", "fisico", "retail", "offline", "brick", "shop",
             "punto_vendita", "pv", "instore", "in_store", "fisica"}


class InputError(ValueError):
    """A problem with the user's file, with a message meant for the user."""


# --------------------------------------------------------------------------- reading
def file_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _as_bytes(source: str | Path | bytes | BinaryIO) -> bytes:
    if isinstance(source, (str, Path)):
        return Path(source).read_bytes()
    if isinstance(source, (bytes, bytearray)):
        return bytes(source)
    pos = source.tell() if hasattr(source, "tell") else None
    data = source.read()
    if pos is not None and hasattr(source, "seek"):
        source.seek(pos)
    return data


def _extension(source, filename: str | None) -> str:
    name = filename or getattr(source, "name", None) or (str(source) if isinstance(source, (str, Path)) else "")
    ext = Path(str(name)).suffix.lower()
    if not ext:
        raise InputError("Cannot tell the file type: use a .csv, .xlsx or .parquet file.")
    return ext


def _decode_sample(data: bytes) -> tuple[str, str]:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data[:65536].decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise InputError("The file is not readable text (unknown encoding).")


def _sniff_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        first_line = sample.splitlines()[0] if sample else ""
        counts = {d: first_line.count(d) for d in ",;\t|"}
        best = max(counts, key=counts.get)
        return best if counts[best] > 0 else ","


def read_any(source: str | Path | bytes | BinaryIO, filename: str | None = None,
             sheet: str | int | None = None) -> pd.DataFrame:
    """Read a ``.csv`` (delimiter and encoding sniffed), ``.xlsx`` or ``.parquet`` file.

    CSV columns are read as text and converted later by ``standardise`` (this keeps
    leading zeros in codes and handles decimal commas). Raises ``InputError``.
    """
    ext = _extension(source, filename)
    try:
        if ext in (".csv", ".txt", ".tsv"):
            data = _as_bytes(source)
            sample, encoding = _decode_sample(data)
            sep = "\t" if ext == ".tsv" else _sniff_delimiter(sample)
            df = pd.read_csv(io.BytesIO(data), sep=sep, dtype=str, encoding=encoding,
                             skipinitialspace=True)
        elif ext in (".xlsx", ".xlsm"):
            df = pd.read_excel(io.BytesIO(_as_bytes(source)), sheet_name=0 if sheet is None else sheet,
                               engine="openpyxl")
        elif ext == ".xls":
            raise InputError("Old .xls files are not supported: save the file as .xlsx or .csv.")
        elif ext in (".parquet", ".pq"):
            df = pd.read_parquet(io.BytesIO(_as_bytes(source)))
        else:
            raise InputError(f"Unsupported file type {ext!r}: use .csv, .xlsx or .parquet.")
    except InputError:
        raise
    except Exception as exc:  # pandas / openpyxl / pyarrow errors
        raise InputError(f"Could not read the file: {exc}") from exc
    df.columns = [str(c).strip() for c in df.columns]
    if df.empty or len(df.columns) == 0:
        raise InputError("The file has no rows.")
    return df


def list_sheets(source: str | Path | bytes | BinaryIO) -> list[str]:
    """Sheet names of an Excel workbook (for the sheet picker)."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(_as_bytes(source)), read_only=True)
    except Exception as exc:
        raise InputError(f"Could not open the workbook: {exc}") from exc
    return list(wb.sheetnames)


# --------------------------------------------------------------------------- mapping
def normalise_name(name: str) -> str:
    text = unicodedata.normalize("NFKD", str(name))
    text = "".join(ch for ch in text if not unicodedata.combining(ch)).lower()
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


def guess_mapping(columns) -> dict[str, str | None]:
    """Standard field -> user column, guessed from column names (EN/IT synonyms).

    Exact synonym matches win over partial (token) matches; each column is used
    at most once. Unmatched fields map to ``None``.
    """
    cols = [str(c) for c in columns]
    candidates = []
    for c_pos, col in enumerate(cols):
        norm = normalise_name(col)
        tokens = set(norm.split("_"))
        for f_pos, fld in enumerate(FIELDS):
            syns = SYNONYMS[fld]
            if norm in syns:
                candidates.append((3, -syns.index(norm), -f_pos, -c_pos, fld, col))
            elif norm.replace("_", "") in {s.replace("_", "") for s in syns}:
                candidates.append((2, 0, -f_pos, -c_pos, fld, col))
            else:
                hits = [s for s in syns if "_" not in s and len(s) > 2 and s in tokens]
                if hits:
                    candidates.append((1, -syns.index(hits[0]), -f_pos, -c_pos, fld, col))
    mapping: dict[str, str | None] = {f: None for f in FIELDS}
    used: set[str] = set()
    for *_, fld, col in sorted(candidates, reverse=True):
        if mapping[fld] is None and col not in used:
            mapping[fld] = col
            used.add(col)
    return mapping


@dataclass(frozen=True)
class ColumnMapping:
    date: str
    item: str
    quantity: str
    store: str | None = None
    unit_price: str | None = None
    unit_cost: str | None = None
    channel: str | None = None
    lead_time_days: str | None = None

    @classmethod
    def from_dict(cls, mapping: dict[str, str | None]) -> ColumnMapping:
        missing = [f for f in REQUIRED if not mapping.get(f)]
        if missing:
            raise InputError(f"Map the required column(s): {', '.join(missing)}.")
        return cls(**{f: (mapping.get(f) or None) for f in FIELDS})

    def as_dict(self) -> dict[str, str | None]:
        return {f: getattr(self, f) for f in FIELDS}


# --------------------------------------------------------------------------- parsing
_ISO = re.compile(r"^\d{4}-\d{1,2}-\d{1,2}")
_DMY = re.compile(r"^(\d{1,2})([/.\-])(\d{1,2})\2(\d{2}|\d{4})$")
_YMD8 = re.compile(r"^\d{8}$")


def parse_dates(values: pd.Series, fmt: str = "auto") -> tuple[pd.Series, list[str]]:
    """Parse a column of dates; returns (datetime64 floored to day, notes)."""
    notes: list[str] = []
    if pd.api.types.is_datetime64_any_dtype(values):
        parsed = pd.to_datetime(values)
        if getattr(parsed.dt, "tz", None) is not None:
            parsed = parsed.dt.tz_localize(None)
        return parsed.dt.floor("D"), notes
    if pd.api.types.is_numeric_dtype(values):
        num = pd.to_numeric(values, errors="coerce")
        if num.dropna().between(19000101, 21001231).all():
            return pd.to_datetime(num.astype("Int64").astype(str), format="%Y%m%d", errors="coerce"), notes
        notes.append("Dates stored as Excel serial numbers were converted.")
        return pd.to_datetime(num, unit="D", origin="1899-12-30", errors="coerce").dt.floor("D"), notes

    text = values.astype("string").str.strip()
    date_part = text.str.split(r"[ T]", n=1, regex=True).str[0]
    sample = date_part.dropna()
    sample = sample[sample != ""].head(5000)
    if sample.empty:
        return pd.Series(pd.NaT, index=values.index), notes

    if fmt == "ISO" or (fmt == "auto" and sample.str.match(_ISO).mean() > 0.9):
        parsed = pd.to_datetime(date_part, format="%Y-%m-%d", errors="coerce")
    elif sample.str.match(_YMD8).mean() > 0.9:
        parsed = pd.to_datetime(date_part, format="%Y%m%d", errors="coerce")
    elif sample.str.match(_DMY).mean() > 0.9:
        parts = sample.str.extract(_DMY)
        first, second = parts[0].astype(int), parts[2].astype(int)
        sep = parts[1].mode().iloc[0]
        year = "%Y" if parts[3].str.len().max() == 4 else "%y"
        if fmt == "dayfirst":
            dayfirst = True
        elif fmt == "monthfirst":
            dayfirst = False
        elif (first > 12).any():
            dayfirst = True
        elif (second > 12).any():
            dayfirst = False
        else:
            dayfirst = True
            notes.append("Dates like 03/04/2024 are ambiguous: day-first was assumed (set the date format to change it).")
        pattern = f"%d{sep}%m{sep}{year}" if dayfirst else f"%m{sep}%d{sep}{year}"
        parsed = pd.to_datetime(date_part, format=pattern, errors="coerce")
    else:
        parsed = pd.to_datetime(date_part, errors="coerce", dayfirst=fmt != "monthfirst", format="mixed")
    return parsed.dt.floor("D"), notes


_THOUSANDS = re.compile(r"^-?\d{1,3}(,\d{3})+$")


def parse_numbers(values: pd.Series) -> pd.Series:
    """Parse numbers written as ``12``, ``12.5``, ``12,5``, ``1.234,5`` or ``1,234.5``."""
    if pd.api.types.is_numeric_dtype(values):
        return pd.to_numeric(values, errors="coerce").astype("float64")
    s = values.astype("string").str.strip().str.replace(r"[€$£\s']", "", regex=True)
    has_comma = s.str.contains(",", regex=False).fillna(False)
    has_dot = s.str.contains(".", regex=False).fillna(False)
    both = has_comma & has_dot
    last_comma = s.str.rfind(",") > s.str.rfind(".")
    s = s.mask(both & last_comma, s.str.replace(".", "", regex=False).str.replace(",", ".", regex=False))
    s = s.mask(both & ~last_comma, s.str.replace(",", "", regex=False))
    comma_only = has_comma & ~has_dot
    if comma_only.any():
        vals = s[comma_only].dropna()
        thousands = len(vals) > 0 and vals.str.match(_THOUSANDS).all()
        repl = s.str.replace(",", "", regex=False) if thousands else s.str.replace(",", ".", regex=False)
        s = s.mask(comma_only, repl)
    return pd.to_numeric(s, errors="coerce").astype("float64")


def normalise_channel(values: pd.Series) -> tuple[pd.Series, int]:
    """Map free-text channel labels to physical/online; returns (labels, n_unknown)."""
    norm = values.astype("string").map(lambda v: normalise_name(v) if isinstance(v, str) else "")
    out = pd.Series("physical", index=values.index, dtype="object")
    online = norm.isin(_ONLINE)
    out[online] = "online"
    unknown = ~online & ~norm.isin(_PHYSICAL) & (norm != "")
    return out, int(unknown.sum())


# --------------------------------------------------------------------------- standardise
@dataclass
class StandardisedData:
    sales: pd.DataFrame          # [date, store, item, sales] zero-filled from each series' first day
    item_master: pd.DataFrame    # [item, unit_price, unit_cost, lead_time_days, price_source, cost_source]
    stores: pd.DataFrame         # [store, channel, channel_source]
    source: str                  # "demo" | "file"
    file_name: str = ""
    file_hash: str = ""
    notes: dict = field(default_factory=dict)

    @property
    def date_min(self) -> pd.Timestamp:
        return self.sales["date"].min()

    @property
    def date_max(self) -> pd.Timestamp:
        return self.sales["date"].max()

    @property
    def n_series(self) -> int:
        return int(self.sales[["store", "item"]].drop_duplicates().shape[0])

    def matrix(self) -> DemandMatrix:
        return DemandMatrix.from_long(self.sales)


def standardise(raw: pd.DataFrame, mapping: ColumnMapping, cfg: InputConfig, prices: PricesConfig,
                inventory: InventoryParams, source: str = "file", file_name: str = "",
                file_hash_value: str = "") -> StandardisedData:
    """Apply the input contract. Raises ``InputError`` on blocking problems.

    1) rename columns  2) parse dates and numbers  3) IDs as text  4) negative
    quantities -> clip or reject  5) sum per (date, store, item)  6) zero-fill
    each series from its first day to the last day in the file  7) item master
    (last known price / cost / lead time) and store table (channel), with
    documented fallbacks flagged as *assumed*.
    """
    for fld, col in mapping.as_dict().items():
        if col is not None and col not in raw.columns:
            raise InputError(f"Column {col!r} (mapped to {fld}) is not in the file.")

    notes: dict = {"n_rows_raw": int(len(raw)), "warnings": [], "assumptions": []}
    warnings, assumptions = notes["warnings"], notes["assumptions"]

    dates, date_notes = parse_dates(raw[mapping.date], cfg.date_format)
    warnings.extend(date_notes)
    qty = parse_numbers(raw[mapping.quantity])
    item = raw[mapping.item].astype("string").str.strip()
    if mapping.store:
        store = raw[mapping.store].astype("string").str.strip().fillna("")
        store = store.mask(store == "", SINGLE_STORE)
    else:
        store = pd.Series(SINGLE_STORE, index=raw.index, dtype="string")
        assumptions.append("No store column: all rows are treated as one location (store = ALL).")

    bad_date = dates.isna()
    bad_item = item.isna() | (item == "")
    bad_qty = qty.isna()
    if bad_date.all():
        raise InputError(f"No date in column {mapping.date!r} could be read.")
    if bad_qty.all():
        raise InputError(f"No number in column {mapping.quantity!r} could be read.")
    drop = bad_date | bad_item | bad_qty
    for mask, label in ((bad_date, "unreadable dates"), (bad_item, "an empty item code"),
                        (bad_qty & ~bad_date, "an unreadable quantity")):
        if mask.any():
            warnings.append(f"{int(mask.sum()):,} rows dropped because of {label}.")
    negative = (qty < 0) & ~drop
    if negative.any():
        if cfg.negative_quantity == "reject":
            raise InputError(f"{int(negative.sum()):,} rows have a negative quantity (returns).")
        qty = qty.where(~negative, 0.0)
        warnings.append(f"{int(negative.sum()):,} negative quantities (returns) were set to 0.")
    notes["n_rows_dropped"] = int(drop.sum())

    keep = ~drop
    df = pd.DataFrame({"date": dates[keep], "store": store[keep].astype(str),
                       "item": item[keep].astype(str), "sales": qty[keep]})
    if df.empty:
        raise InputError("No usable rows after reading dates, items and quantities.")
    n_before = len(df)
    df = df.groupby(["store", "item", "date"], as_index=False, sort=False)["sales"].sum()
    if len(df) < n_before:
        warnings.append(f"{n_before - len(df):,} rows with the same day, store and item were summed.")

    matrix = DemandMatrix.from_long(df)
    sales = matrix.to_long()
    n_filled = len(sales) - len(df)
    if n_filled > 0:
        notes["n_days_zero_filled"] = int(n_filled)

    # ---- item master
    attrs = pd.DataFrame({"date": dates[keep], "item": item[keep].astype(str)})
    for fld in ("unit_price", "unit_cost", "lead_time_days"):
        col = getattr(mapping, fld)
        attrs[fld] = parse_numbers(raw.loc[keep, col]) if col else np.nan
    attrs.loc[attrs["unit_price"] <= 0, "unit_price"] = np.nan
    attrs.loc[attrs["unit_cost"] <= 0, "unit_cost"] = np.nan
    attrs.loc[attrs["lead_time_days"] < 1, "lead_time_days"] = np.nan
    master = attrs.sort_values("date").groupby("item")[["unit_price", "unit_cost", "lead_time_days"]].last()
    items = natural_sorted(sales["item"].unique())
    master = master.reindex(items)
    master.index.name = "item"
    master = _fill_item_master(master.reset_index(), prices, inventory)

    # ---- stores
    store_ids = natural_sorted(sales["store"].unique())
    if mapping.channel:
        ch, n_unknown = normalise_channel(raw.loc[keep, mapping.channel])
        tmp = pd.DataFrame({"store": store[keep].astype(str), "channel": ch})
        stores = (tmp.groupby("store")["channel"]
                  .agg(lambda s: s.mode().iloc[0]).reindex(store_ids).fillna("physical").reset_index())
        mixed = tmp.groupby("store")["channel"].nunique()
        if (mixed > 1).any():
            warnings.append(f"{int((mixed > 1).sum())} store(s) have more than one channel: the most frequent was used.")
        if n_unknown:
            warnings.append(f"{n_unknown:,} rows have an unknown channel label and were read as physical.")
        stores["channel_source"] = "file"
    else:
        stores = pd.DataFrame({"store": store_ids, "channel": "physical", "channel_source": "default"})

    return StandardisedData(sales=sales, item_master=master, stores=stores, source=source,
                            file_name=file_name, file_hash=file_hash_value, notes=notes)


def _fill_item_master(master: pd.DataFrame, prices: PricesConfig, inventory: InventoryParams) -> pd.DataFrame:
    """Fill missing prices (synthetic, seeded) and costs (ratio x price); record the source."""
    from src.data import make_price_table

    master = master.copy()
    if "price_source" not in master:
        master["price_source"] = np.where(master["unit_price"].notna(), "file", "assumed")
    if "cost_source" not in master:
        master["cost_source"] = np.where(master["unit_cost"].notna(), "file", "derived")
    missing_price = master["unit_price"].isna()
    if missing_price.any():
        synthetic = make_price_table(master.loc[missing_price, "item"].tolist(), prices,
                                     inventory.cost_to_price_ratio).set_index("item")["unit_price"]
        master.loc[missing_price, "unit_price"] = master.loc[missing_price, "item"].map(synthetic).to_numpy()
        master.loc[missing_price, "price_source"] = "assumed"
    missing_cost = master["unit_cost"].isna() | (master["cost_source"] == "derived")
    master.loc[missing_cost, "unit_cost"] = master.loc[missing_cost, "unit_price"] * inventory.cost_to_price_ratio
    master.loc[missing_cost, "cost_source"] = "derived"
    master["lead_time_days"] = pd.to_numeric(master["lead_time_days"], errors="coerce").round().astype("Int64")
    return master[["item", "unit_price", "unit_cost", "lead_time_days", "price_source", "cost_source"]]


def item_assumptions(master: pd.DataFrame, prices: PricesConfig, inventory: InventoryParams) -> list[str]:
    """Human-readable list of the item-level values that were assumed."""
    out = []
    n_price = int((master["price_source"] == "assumed").sum())
    n_cost = int((master["cost_source"] == "derived").sum())
    n_lt = int(master["lead_time_days"].isna().sum())
    if n_price:
        out.append(f"Selling price missing for {n_price} item(s): synthetic prices "
                   f"(€{prices.min:.0f}–€{prices.max:.0f}, seed {prices.seed}) were assumed.")
    if n_cost:
        out.append(f"Unit cost missing for {n_cost} item(s): cost = "
                   f"{inventory.cost_to_price_ratio:.0%} of the selling price was assumed.")
    if n_lt:
        out.append(f"Lead time missing for {n_lt} item(s): the default of "
                   f"{inventory.lead_time_days} days is used.")
    return out


def merge_item_master(data: StandardisedData, master_raw: pd.DataFrame, prices: PricesConfig,
                      inventory: InventoryParams, mapping: dict[str, str | None] | None = None) -> StandardisedData:
    """Override price / cost / lead time with an item-master file (FR-53)."""
    mapping = mapping or guess_mapping(master_raw.columns)
    if not mapping.get("item"):
        raise InputError("The item master needs an item column.")
    upd = pd.DataFrame({"item": master_raw[mapping["item"]].astype("string").str.strip().astype(str)})
    for fld in ("unit_price", "unit_cost", "lead_time_days"):
        upd[fld] = parse_numbers(master_raw[mapping[fld]]) if mapping.get(fld) else np.nan
    upd = upd.drop_duplicates("item", keep="last").set_index("item")

    master = data.item_master.set_index("item").copy()
    common = master.index.intersection(upd.index)
    for fld, source_col, source in (("unit_price", "price_source", "item_master"),
                                    ("unit_cost", "cost_source", "item_master"),
                                    ("lead_time_days", None, None)):
        new = upd.loc[common, fld]
        new = new[new.notna() & (new > 0)]
        if len(new):
            master.loc[new.index, fld] = new.round().astype("Int64") if fld == "lead_time_days" else new
            if source_col:
                master.loc[new.index, source_col] = source
    master = _fill_item_master(master.reset_index(), prices, inventory)
    notes = dict(data.notes)
    notes["item_master_items"] = int(len(common))
    return StandardisedData(data.sales, master, data.stores, data.source, data.file_name, data.file_hash, notes)


# --------------------------------------------------------------------------- validate
@dataclass
class ValidationReport:
    n_rows_raw: int
    n_rows: int
    n_stores: int
    n_items: int
    n_series: int
    date_min: date | None
    date_max: date | None
    history_days: int
    pct_zero_days: float
    excluded_series: list[tuple[str, str, str]] = field(default_factory=list)
    assumptions: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors

    def to_dict(self) -> dict:
        d = dict(self.__dict__)
        d["date_min"] = str(self.date_min) if self.date_min else None
        d["date_max"] = str(self.date_max) if self.date_max else None
        d["ok"] = self.ok
        return d

    def to_markdown(self) -> str:
        lines = [
            "# Data validation report",
            "",
            f"**Status:** {'✅ OK' if self.ok else '❌ blocked'}",
            "",
            "| Check | Value |",
            "|---|---|",
            f"| Rows in the file | {self.n_rows_raw:,} |",
            f"| Rows after standardisation (daily, zero-filled) | {self.n_rows:,} |",
            f"| Stores · items · series | {self.n_stores:,} · {self.n_items:,} · {self.n_series:,} |",
            f"| Period | {self.date_min} → {self.date_max} ({self.history_days:,} days) |",
            f"| Days with zero sales | {self.pct_zero_days:.1%} |",
            f"| Series excluded | {len(self.excluded_series):,} |",
        ]
        for title, items in (("Errors", self.errors), ("Warnings", self.warnings),
                             ("Assumptions", self.assumptions)):
            if items:
                lines += ["", f"## {title}", ""] + [f"- {x}" for x in items]
        if self.excluded_series:
            lines += ["", "## Excluded series", "", "| Store | Item | Reason |", "|---|---|---|"]
            lines += [f"| {s} | {i} | {r} |" for s, i, r in self.excluded_series[:200]]
            if len(self.excluded_series) > 200:
                lines.append(f"| … | … | {len(self.excluded_series) - 200} more |")
        return "\n".join(lines) + "\n"


def validate(data: StandardisedData, cfg: InputConfig, enforce_limits: bool = False,
             prices: PricesConfig | None = None, inventory: InventoryParams | None = None) -> ValidationReport:
    """Check the standardised data. Never raises: errors block the run, warnings don't."""
    sales = data.sales
    n_series = data.n_series
    dmin, dmax = data.date_min, data.date_max
    history = int((dmax - dmin).days) + 1 if len(sales) else 0
    report = ValidationReport(
        n_rows_raw=int(data.notes.get("n_rows_raw", len(sales))),
        n_rows=int(len(sales)),
        n_stores=int(sales["store"].nunique()),
        n_items=int(sales["item"].nunique()),
        n_series=n_series,
        date_min=dmin.date() if len(sales) else None,
        date_max=dmax.date() if len(sales) else None,
        history_days=history,
        pct_zero_days=float((sales["sales"] == 0).mean()) if len(sales) else 0.0,
        assumptions=list(data.notes.get("assumptions", []))
        + item_assumptions(data.item_master, prices or PricesConfig(), inventory or InventoryParams()),
        warnings=list(data.notes.get("warnings", [])),
    )
    if history < cfg.min_history_days:
        report.errors.append(
            f"Only {history} days of history: at least {cfg.min_history_days} days (8 weeks) are needed.")
    if n_series == 0:
        report.errors.append("No series left after reading the file.")
    if enforce_limits:
        lim = cfg.limits
        if report.n_rows_raw > lim.max_rows:
            report.errors.append(
                f"{report.n_rows_raw:,} rows is above the online limit of {lim.max_rows:,}: run the analysis locally (see README).")
        if n_series > lim.max_series:
            report.errors.append(
                f"{n_series:,} store-item series is above the online limit of {lim.max_series:,}: run the analysis locally (see README).")
    if report.pct_zero_days > 0.5:
        report.warnings.append(
            f"{report.pct_zero_days:.0%} of the days have zero sales (intermittent demand): forecasts will be less reliable.")
    if 0 < history < 400 + 28:
        report.warnings.append("Less than about 14 months of history: yearly seasonality will not be modelled.")
    return report
